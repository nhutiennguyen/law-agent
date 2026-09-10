# backend/core/contract_parser.py — Module trích xuất & xử lý văn bản hợp đồng

import io
from typing import Tuple
from pypdf import PdfReader
from docx import Document

class ContractParser:
    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, bool]:
        """
        Trích xuất văn bản từ file bytes.
        Trả về: (text_content, is_scanned_or_image)
        """
        filename_lower = filename.lower()

        # 1. Định dạng TXT / Markdown
        if filename_lower.endswith(".txt") or filename_lower.endswith(".md"):
            try:
                return file_bytes.decode("utf-8"), False
            except UnicodeDecodeError:
                try:
                    return file_bytes.decode("cp1258"), False  # Vietnamese Windows encoding
                except Exception:
                    return file_bytes.decode("latin-1", errors="ignore"), False

        # 2. Định dạng PDF
        elif filename_lower.endswith(".pdf"):
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                extracted_pages = []
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        extracted_pages.append(f"--- TRANG {i + 1} ---\n{text}")
                
                full_text = "\n\n".join(extracted_pages).strip()
                # Nếu file PDF chứa rất ít text (< 50 từ), nhiều khả năng là file scan dạng ảnh
                if len(full_text.split()) < 50:
                    return full_text, True
                return full_text, False
            except Exception as e:
                raise ValueError(f"Không thể đọc file PDF: {str(e)}")

        # 3. Định dạng DOCX (Microsoft Word)
        elif filename_lower.endswith(".docx"):
            try:
                doc = Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Đọc cả bảng biểu trong Word
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                return "\n".join(paragraphs), False
            except Exception as e:
                raise ValueError(f"Không thể đọc file DOCX: {str(e)}")

        # 4. Định dạng Hình ảnh (JPG, PNG, WebP)
        elif any(filename_lower.endswith(ext) for ext in [".jpg", ".jpeg", ".png", ".webp"]):
            return "", True

        else:
            raise ValueError(
                f"Định dạng file '{filename}' không được hỗ trợ. "
                "Hệ thống hỗ trợ file .PDF, .DOCX, .TXT, .PNG, .JPG."
            )

contract_parser = ContractParser()
